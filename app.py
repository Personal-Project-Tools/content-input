from __future__ import annotations

import csv
import hashlib
import html
import io
import json
import re
import sys
import traceback
import zipfile
from dataclasses import dataclass, field
from datetime import datetime
from email.parser import BytesParser
from email.policy import default
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
from typing import Iterable
from xml.etree.ElementTree import iterparse

try:
    from openpyxl import Workbook, load_workbook
except ImportError as exc:  # pragma: no cover - shown to local user at startup
    raise SystemExit(
        "Missing dependency: openpyxl. Install it with: python -m pip install -r requirements.txt"
    ) from exc


APP_NAME = "TTK Content Tool"
DEFAULT_PORT = 8765
OUTPUT_DIR = Path("outputs")

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")
if hasattr(sys.stderr, "reconfigure"):
    sys.stderr.reconfigure(encoding="utf-8")


@dataclass(frozen=True)
class MetricSpec:
    label: str
    aliases: tuple[str, ...]
    fallback_col: int


@dataclass(frozen=True)
class ContentSpec:
    label: str
    filename_markers: tuple[str, ...]
    header_markers: tuple[str, ...]
    date_aliases: tuple[str, ...]
    date_fallback_col: int
    metrics: dict[str, MetricSpec]


CONTENT_SPECS = {
    "Live": ContentSpec(
        label="Live",
        filename_markers=("live",),
        header_markers=("LIVE-attributed GMV", "Launched Time", "Duration"),
        date_aliases=("Launched Time",),
        date_fallback_col=4,
        metrics={
            "revenue_24h": MetricSpec(
                "revenue_24h", ("LIVE-attributed GMV",), 6
            ),
            "content_impression": MetricSpec("content_impression", ("Views",), 20),
            "product_impression": MetricSpec(
                "product_impression", ("Product Impressions",), 26
            ),
        },
    ),
    "Video": ContentSpec(
        label="Video",
        filename_markers=("video",),
        header_markers=("Video Info", "Video ID", "Video Views", "VV", "Video-attributed GMV"),
        date_aliases=("Time",),
        date_fallback_col=5,
        metrics={
            "revenue_24h": MetricSpec(
                "revenue_24h", ("Video-attributed GMV",), 22
            ),
            "content_impression": MetricSpec("content_impression", ("Video Views", "Video views", "VV"), 7),
            "product_impression": MetricSpec(
                "product_impression", ("Product Impressions",), 13
            ),
        },
    ),
}


@dataclass
class ParsedFile:
    filename: str
    content_type: str
    period: str
    revenue_24h: float
    content_impression: float
    product_impression: float
    rows: int
    daily_rows: list["SummaryRow"] = field(default_factory=list)
    digest: str = ""
    warnings: list[str] = field(default_factory=list)


@dataclass
class SummaryRow:
    content_type: str
    period: str
    revenue_24h: float = 0.0
    content_impression: float = 0.0
    product_impression: float = 0.0
    files: list[str] = field(default_factory=list)
    rows: int = 0
    warnings: list[str] = field(default_factory=list)


def normalize_text(value: object) -> str:
    if value is None:
        return ""
    return re.sub(r"\s+", " ", str(value).strip()).casefold()


def to_number(value: object) -> float:
    if value is None or value == "":
        return 0.0
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip()
    if not text:
        return 0.0
    text = re.sub(r"[^\d.\-]", "", text.replace(",", ""))
    if text in {"", "-", "."}:
        return 0.0
    try:
        return float(text)
    except ValueError:
        return 0.0


def extract_period(a1_value: object) -> str:
    text = str(a1_value or "")
    match = re.search(r"(\d{4})-(\d{2})-\d{2}", text)
    if match:
        return f"Tháng {match.group(2)}/{match.group(1)}"
    match = re.search(r"(\d{4})/(\d{2})/\d{2}", text)
    if match:
        return f"Tháng {match.group(2)}/{match.group(1)}"
    return "Không rõ"


def extract_day(value: object) -> str:
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d")
    text = str(value or "").strip()
    if not text:
        return "KhÃ´ng rÃµ"
    match = re.search(r"(\d{4})[/-](\d{1,2})[/-](\d{1,2})", text)
    if not match:
        return "KhÃ´ng rÃµ"
    year, month, day = (int(part) for part in match.groups())
    try:
        return datetime(year, month, day).strftime("%Y-%m-%d")
    except ValueError:
        return "KhÃ´ng rÃµ"


def detect_header_row(ws) -> int:
    best_row = 1
    best_score = -1
    limit = min(ws.max_row or 1, 20)
    for row_idx in range(1, limit + 1):
        values = [normalize_text(cell.value) for cell in ws[row_idx]]
        joined = " | ".join(values)
        score = 0
        for spec in CONTENT_SPECS.values():
            score += sum(1 for marker in spec.header_markers if marker.casefold() in joined)
            for metric in spec.metrics.values():
                score += sum(1 for alias in metric.aliases if alias.casefold() in joined)
        if score > best_score:
            best_score = score
            best_row = row_idx
    return best_row


def detect_content_type(filename: str, header_values: Iterable[object]) -> str:
    lower_name = filename.casefold()
    for name, spec in CONTENT_SPECS.items():
        if any(marker in lower_name for marker in spec.filename_markers):
            return name

    joined = " | ".join(normalize_text(value) for value in header_values)
    scores = {
        name: sum(1 for marker in spec.header_markers if marker.casefold() in joined)
        for name, spec in CONTENT_SPECS.items()
    }
    winner, score = max(scores.items(), key=lambda item: item[1])
    if score <= 0:
        raise ValueError("Không nhận diện được file Live hay Video.")
    return winner


def find_metric_columns(ws, header_row: int, spec: ContentSpec) -> tuple[dict[str, int], list[str]]:
    header_values = [ws.cell(header_row, col).value for col in range(1, ws.max_column + 1)]
    normalized = [normalize_text(value) for value in header_values]
    columns: dict[str, int] = {}
    warnings: list[str] = []

    for metric_name, metric in spec.metrics.items():
        found_col = None
        for col_idx, header in enumerate(normalized, start=1):
            if any(alias.casefold() == header or alias.casefold() in header for alias in metric.aliases):
                found_col = col_idx
                break
        if found_col is None:
            found_col = metric.fallback_col
            warnings.append(
                f"Không thấy header cho {metric.label}; dùng cột fallback {column_letter(found_col)}."
            )
        columns[metric_name] = found_col

    return columns, warnings


def find_date_column(ws, header_row: int, spec: ContentSpec) -> tuple[int, list[str]]:
    header_values = [ws.cell(header_row, col).value for col in range(1, ws.max_column + 1)]
    normalized = [normalize_text(value) for value in header_values]

    for col_idx, header in enumerate(normalized, start=1):
        if any(alias.casefold() in header for alias in spec.date_aliases):
            return col_idx, []

    return spec.date_fallback_col, [
        f"KhÃ´ng tháº¥y header ngÃ y; dÃ¹ng cá»™t fallback {column_letter(spec.date_fallback_col)}."
    ]


def add_daily_totals(
    grouped: dict[tuple[str, str], SummaryRow],
    content_type: str,
    day: str,
    filename: str,
    totals: dict[str, float],
) -> None:
    row = grouped.setdefault((content_type, day), SummaryRow(content_type, day))
    row.revenue_24h += totals["revenue_24h"]
    row.content_impression += totals["content_impression"]
    row.product_impression += totals["product_impression"]
    if filename not in row.files:
        row.files.append(filename)
    row.rows += 1


def column_letter(index: int) -> str:
    letters = ""
    while index:
        index, rem = divmod(index - 1, 26)
        letters = chr(65 + rem) + letters
    return letters


def column_index(ref: str) -> int:
    letters = re.match(r"[A-Z]+", ref.upper())
    if not letters:
        return 0
    index = 0
    for char in letters.group(0):
        index = index * 26 + ord(char) - 64
    return index


def row_index(ref: str) -> int:
    match = re.search(r"\d+", ref)
    return int(match.group(0)) if match else 0


def read_shared_strings(zf: zipfile.ZipFile) -> list[str]:
    try:
        handle = zf.open("xl/sharedStrings.xml")
    except KeyError:
        return []

    strings: list[str] = []
    ns = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
    with handle:
        for _, elem in iterparse(handle, events=("end",)):
            if elem.tag == f"{ns}si":
                parts = [node.text or "" for node in elem.iter(f"{ns}t")]
                strings.append("".join(parts))
                elem.clear()
    return strings


def read_shared_string_at(zf: zipfile.ZipFile, target_index: int) -> str:
    try:
        handle = zf.open("xl/sharedStrings.xml")
    except KeyError:
        return ""

    ns = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
    current_index = -1
    with handle:
        for _, elem in iterparse(handle, events=("end",)):
            if elem.tag == f"{ns}si":
                current_index += 1
                if current_index == target_index:
                    parts = [node.text or "" for node in elem.iter(f"{ns}t")]
                    return "".join(parts)
                elem.clear()
    return ""


def first_sheet_name(zf: zipfile.ZipFile) -> str:
    sheet_names = sorted(
        name
        for name in zf.namelist()
        if name.startswith("xl/worksheets/sheet") and name.endswith(".xml")
    )
    if not sheet_names:
        raise ValueError("Không tìm thấy worksheet trong file Excel.")
    return sheet_names[0]


def cell_text(cell, shared_strings: list[str] | None, ns: str) -> str:
    cell_type = cell.attrib.get("t")
    if cell_type == "inlineStr":
        return "".join(node.text or "" for node in cell.iter(f"{ns}t"))

    value_node = cell.find(f"{ns}v")
    if value_node is None or value_node.text is None:
        return ""
    raw = value_node.text
    if cell_type == "s":
        if shared_strings is None:
            return ""
        try:
            return shared_strings[int(raw)]
        except (ValueError, IndexError):
            return ""
    return raw


def cell_raw_value(cell, ns: str) -> str:
    value_node = cell.find(f"{ns}v")
    return value_node.text if value_node is not None and value_node.text is not None else ""


def detect_content_type_from_filename(filename: str) -> str | None:
    lower_name = filename.casefold()
    for name, spec in CONTENT_SPECS.items():
        if any(marker in lower_name for marker in spec.filename_markers):
            return name
    return None


def parse_workbook_fast(file_bytes: bytes, filename: str) -> ParsedFile:
    content_type = detect_content_type_from_filename(filename)
    if content_type is None:
        raise ValueError("Fast parser cần tên file có Live hoặc Video.")

    spec = CONTENT_SPECS[content_type]
    target_cols = {metric.fallback_col: name for name, metric in spec.metrics.items()}
    date_col = spec.date_fallback_col
    totals = {"revenue_24h": 0.0, "content_impression": 0.0, "product_impression": 0.0}
    daily_grouped: dict[tuple[str, str], SummaryRow] = {}
    period = "Không rõ"
    day = "Không rõ"
    a1_shared_index: int | None = None
    data_rows = 0
    ns = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"

    with zipfile.ZipFile(io.BytesIO(file_bytes)) as zf:
        shared_strings = read_shared_strings(zf)
        sheet_name = first_sheet_name(zf)
        current_row = 0
        current_day = "KhÃ´ng rÃµ"
        current_totals = {
            "revenue_24h": 0.0,
            "content_impression": 0.0,
            "product_impression": 0.0,
        }
        current_has_values = False
        header_detected = False
        row_cells: dict[int, str] = {}
        with zf.open(sheet_name) as sheet:
            for _, elem in iterparse(sheet, events=("end",)):
                if elem.tag == f"{ns}row":
                    row_num = int(elem.attrib.get("r", "0") or 0)
                    if not header_detected and 1 <= row_num <= 5:
                        matched = {}
                        for metric_name, metric in spec.metrics.items():
                            for c_idx, text in row_cells.items():
                                norm = normalize_text(text)
                                if any(alias.casefold() == norm or alias.casefold() in norm for alias in metric.aliases):
                                    matched[c_idx] = metric_name
                                    break
                        if len(matched) >= 2:
                            header_detected = True
                            target_cols = matched
                            for c_idx, text in row_cells.items():
                                norm = normalize_text(text)
                                if any(alias.casefold() in norm for alias in spec.date_aliases):
                                    date_col = c_idx
                                    break
                    row_cells.clear()

                    if row_num >= 4 and row_num == current_row and current_has_values:
                        data_rows += 1
                        add_daily_totals(
                            daily_grouped,
                            content_type,
                            current_day,
                            filename,
                            current_totals,
                        )
                    current_row = 0
                    current_day = "Không rõ"
                    current_totals = {
                        "revenue_24h": 0.0,
                        "content_impression": 0.0,
                        "product_impression": 0.0,
                    }
                    current_has_values = False
                    elem.clear()
                    continue

                if elem.tag != f"{ns}c":
                    continue

                ref = elem.attrib.get("r", "")
                col_num = column_index(ref)
                row_num = row_index(ref)
                if row_num >= 4 and row_num != current_row:
                    current_row = row_num
                    current_day = "Không rõ"
                    current_totals = {
                        "revenue_24h": 0.0,
                        "content_impression": 0.0,
                        "product_impression": 0.0,
                    }
                    current_has_values = False

                if ref == "A1":
                    if elem.attrib.get("t") == "s":
                        raw_a1 = cell_raw_value(elem, ns)
                        a1_shared_index = int(raw_a1) if raw_a1.isdigit() else None
                    else:
                        a1_text = cell_text(elem, shared_strings, ns)
                        period = extract_period(a1_text)
                        day = extract_day(a1_text)
                elif not header_detected and 1 <= row_num <= 5:
                    t = cell_text(elem, shared_strings, ns)
                    if t:
                        row_cells[col_num] = t
                elif row_num >= 4 and col_num == date_col:
                    current_day = extract_day(cell_text(elem, shared_strings, ns))
                    current_has_values = True
                elif row_num >= 4 and col_num in target_cols:
                    metric_name = target_cols[col_num]
                    value = to_number(cell_raw_value(elem, ns))
                    totals[metric_name] += value
                    current_totals[metric_name] += value
                    current_has_values = True

                elem.clear()

        if period == "Không rõ" and a1_shared_index is not None:
            a1_text = read_shared_string_at(zf, a1_shared_index)
            period = extract_period(a1_text)
            day = extract_day(a1_text)

    return ParsedFile(
        filename=filename,
        content_type=content_type,
        period=period,
        revenue_24h=totals["revenue_24h"],
        content_impression=totals["content_impression"],
        product_impression=totals["product_impression"],
        rows=data_rows,
        daily_rows=[
            SummaryRow(
                content_type,
                day,
                totals["revenue_24h"],
                totals["content_impression"],
                totals["product_impression"],
                [filename],
                data_rows,
            )
        ],
    )


def parse_workbook_slow(file_bytes: bytes, filename: str) -> ParsedFile:
    wb = load_workbook(io.BytesIO(file_bytes), read_only=True, data_only=True)
    try:
        ws = wb[wb.sheetnames[0]]
        if (ws.max_column or 0) <= 1:
            wb.close()
            wb = load_workbook(io.BytesIO(file_bytes), read_only=False, data_only=True)
            ws = wb[wb.sheetnames[0]]
        period = extract_period(ws["A1"].value)
        day = extract_day(ws["A1"].value)
        header_row = detect_header_row(ws)
        header_values = [ws.cell(header_row, col).value for col in range(1, ws.max_column + 1)]
        content_type = detect_content_type(filename, header_values)
        spec = CONTENT_SPECS[content_type]
        columns, warnings = find_metric_columns(ws, header_row, spec)

        totals = {"revenue_24h": 0.0, "content_impression": 0.0, "product_impression": 0.0}
        data_rows = 0
        for row in ws.iter_rows(min_row=header_row + 1, values_only=True):
            if not any(value not in (None, "") for value in row):
                continue
            data_rows += 1
            for metric_name, col_idx in columns.items():
                value = row[col_idx - 1] if col_idx - 1 < len(row) else None
                numeric_value = to_number(value)
                totals[metric_name] += numeric_value

        return ParsedFile(
            filename=filename,
            content_type=content_type,
            period=period,
            revenue_24h=totals["revenue_24h"],
            content_impression=totals["content_impression"],
            product_impression=totals["product_impression"],
            rows=data_rows,
            daily_rows=[
                SummaryRow(
                    content_type,
                    day,
                    totals["revenue_24h"],
                    totals["content_impression"],
                    totals["product_impression"],
                    [filename],
                    data_rows,
                )
            ],
            warnings=warnings,
        )
    finally:
        wb.close()


def parse_workbook(file_bytes: bytes, filename: str) -> ParsedFile:
    try:
        parsed = parse_workbook_fast(file_bytes, filename)
    except Exception:
        parsed = parse_workbook_slow(file_bytes, filename)
    parsed.digest = hashlib.sha256(file_bytes).hexdigest()
    return parsed


def find_duplicate_files(files: Iterable[ParsedFile]) -> list[str]:
    seen: dict[tuple[str, str], str] = {}
    duplicates: list[str] = []
    for parsed in files:
        if not parsed.digest:
            continue
        key = (parsed.content_type, parsed.digest)
        original = seen.get(key)
        if original:
            duplicates.append(
                f"{parsed.filename} giống 100% với {original} trong nhóm {parsed.content_type}."
            )
        else:
            seen[key] = parsed.filename
    return duplicates


def summarize(files: Iterable[ParsedFile]) -> list[SummaryRow]:
    grouped: dict[tuple[str, str], SummaryRow] = {}
    for parsed in files:
        key = (parsed.content_type, parsed.period)
        row = grouped.setdefault(key, SummaryRow(parsed.content_type, parsed.period))
        row.revenue_24h += parsed.revenue_24h
        row.content_impression += parsed.content_impression
        row.product_impression += parsed.product_impression
        row.files.append(parsed.filename)
        row.rows += parsed.rows
        row.warnings.extend(f"{parsed.filename}: {warning}" for warning in parsed.warnings)

    return sorted(grouped.values(), key=lambda item: (item.period, item.content_type))


def summarize_daily(files: Iterable[ParsedFile]) -> list[SummaryRow]:
    grouped: dict[tuple[str, str], SummaryRow] = {}
    for parsed in files:
        for daily in parsed.daily_rows:
            key = (daily.content_type, daily.period)
            row = grouped.setdefault(key, SummaryRow(daily.content_type, daily.period))
            row.revenue_24h += daily.revenue_24h
            row.content_impression += daily.content_impression
            row.product_impression += daily.product_impression
            for filename in daily.files:
                if filename not in row.files:
                    row.files.append(filename)
            row.rows += daily.rows
            row.warnings.extend(daily.warnings)

    return sorted(grouped.values(), key=lambda item: (item.period, item.content_type))


def parse_folder(folder: Path) -> tuple[list[SummaryRow], list[SummaryRow], list[ParsedFile]]:
    parsed_files = []
    for path in sorted(folder.glob("*.xlsx")):
        if path.name.startswith("~$"):
            continue
        parsed_files.append(parse_workbook(path.read_bytes(), path.name))
    return summarize(parsed_files), summarize_daily(parsed_files), parsed_files


def format_number(value: float, decimals: int = 0) -> str:
    rounded = round(value, decimals)
    if decimals == 0:
        return f"{rounded:,.0f}"
    return f"{rounded:,.{decimals}f}"


def summary_to_dict(row: SummaryRow) -> dict[str, object]:
    return {
        "content_type": row.content_type,
        "period": row.period,
        "revenue_24h": row.revenue_24h,
        "content_impression": row.content_impression,
        "product_impression": row.product_impression,
        "file_count": len(row.files),
        "files": row.files,
        "rows": row.rows,
        "warnings": row.warnings,
    }


def write_summary_sheet(ws, headers: list[str], rows: list[SummaryRow]) -> None:
    ws.append(headers)
    for row in rows:
        ws.append(
            [
                row.content_type,
                row.period,
                row.revenue_24h,
                row.product_impression,
                row.content_impression,
                len(row.files),
                "; ".join(row.files),
            ]
        )
    for col in ("C", "D", "E"):
        for cell in ws[col][1:]:
            cell.number_format = "#,##0.00" if col == "C" else "#,##0"
    for col_cells in ws.columns:
        max_len = max(len(str(cell.value or "")) for cell in col_cells)
        ws.column_dimensions[col_cells[0].column_letter].width = min(max(max_len + 2, 12), 60)


def save_outputs(rows: list[SummaryRow], daily_rows: list[SummaryRow] | None = None) -> tuple[Path, Path]:
    daily_rows = daily_rows or []
    OUTPUT_DIR.mkdir(exist_ok=True)
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    csv_path = OUTPUT_DIR / f"ttk_content_summary_{stamp}.csv"
    xlsx_path = OUTPUT_DIR / f"ttk_content_summary_{stamp}.xlsx"

    headers = [
        "Live/Video",
        "Thời gian",
        "revenue_24h",
        "product_impression",
        "content_impression",
        "Số file",
        "Files",
    ]

    with csv_path.open("w", newline="", encoding="utf-8-sig") as handle:
        writer = csv.writer(handle)
        writer.writerow(headers)
        for row in rows:
            writer.writerow(
                [
                    row.content_type,
                    row.period,
                    row.revenue_24h,
                    row.product_impression,
                    row.content_impression,
                    len(row.files),
                    "; ".join(row.files),
                ]
            )

    wb = Workbook()
    ws = wb.active
    ws.title = "TTK_Content"
    write_summary_sheet(ws, headers, rows)
    if daily_rows:
        day_headers = headers.copy()
        day_headers[1] = "Ngay"
        day_ws = wb.create_sheet("TTK_Content_Ngay")
        write_summary_sheet(day_ws, day_headers, daily_rows)
    wb.save(xlsx_path)
    return csv_path, xlsx_path


def render_page(
    error: str = "",
    rows: list[SummaryRow] | None = None,
    daily_rows: list[SummaryRow] | None = None,
) -> bytes:
    rows = rows or []
    daily_rows = daily_rows or []
    table_rows = "\n".join(
        f"""
        <tr>
          <td>{html.escape(row.content_type)}</td>
          <td>{html.escape(row.period)}</td>
          <td class="num">{format_number(row.revenue_24h, 2)}</td>
          <td class="num">{format_number(row.product_impression)}</td>
          <td class="num">{format_number(row.content_impression)}</td>
          <td class="num">{len(row.files)}</td>
          <td>{html.escape("; ".join(row.files))}</td>
        </tr>
        """
        for row in rows
    )
    daily_table_rows = "\n".join(
        f"""
        <tr>
          <td>{html.escape(row.content_type)}</td>
          <td>{html.escape(row.period)}</td>
          <td class="num">{format_number(row.revenue_24h, 2)}</td>
          <td class="num">{format_number(row.product_impression)}</td>
          <td class="num">{format_number(row.content_impression)}</td>
          <td class="num">{len(row.files)}</td>
          <td>{html.escape("; ".join(row.files))}</td>
        </tr>
        """
        for row in daily_rows
    )
    warnings = "\n".join(
        f"<li>{html.escape(warning)}</li>"
        for row in [*rows, *daily_rows]
        for warning in row.warnings
    )
    error_html = f"<div class='error'>{html.escape(error)}</div>" if error else ""
    export_payload = {
        "monthly": [summary_to_dict(row) for row in rows],
        "daily": [summary_to_dict(row) for row in daily_rows],
    }
    table_html = (
        f"""
        <section>
          <div class="actions">
            <form method="post" action="/export">
              <input type="hidden" name="payload" value="{html.escape(json.dumps(export_payload))}">
              <button type="submit">Tải Excel</button>
            </form>
          </div>
          <input class="tab-input" type="radio" name="summary-tab" id="tab-month" checked>
          <input class="tab-input" type="radio" name="summary-tab" id="tab-day">
          <div class="tabs">
            <label for="tab-month">Theo thang</label>
            <label for="tab-day">Theo ngay</label>
          </div>
          <div class="tab-panel month-panel">
          <table>
            <thead>
              <tr>
                <th>Live/Video</th>
                <th>Thời gian</th>
                <th>revenue_24h</th>
                <th>product_impression</th>
                <th>content_impression</th>
                <th>Số file</th>
                <th>Files</th>
              </tr>
            </thead>
            <tbody>{table_rows}</tbody>
          </table>
          </div>
          <div class="tab-panel day-panel">
          <table>
            <thead>
              <tr>
                <th>Live/Video</th>
                <th>Ngay</th>
                <th>revenue_24h</th>
                <th>product_impression</th>
                <th>content_impression</th>
                <th>So file</th>
                <th>Files</th>
              </tr>
            </thead>
            <tbody>{daily_table_rows}</tbody>
          </table>
          </div>
          {"<ul class='warnings'>" + warnings + "</ul>" if warnings else ""}
        </section>
        """
        if rows or daily_rows
        else ""
    )
    return f"""<!doctype html>
<html lang="vi">
<head>
  <meta charset="utf-8">
  <meta name="viewport" content="width=device-width, initial-scale=1">
  <title>{APP_NAME}</title>
  <style>
    :root {{
      color-scheme: light;
      --bg: #f6f7f9;
      --panel: #ffffff;
      --text: #18202a;
      --muted: #667085;
      --line: #d7dde6;
      --accent: #0f766e;
      --accent-dark: #115e59;
      --danger: #b42318;
    }}
    * {{ box-sizing: border-box; }}
    body {{
      margin: 0;
      font: 15px/1.5 Arial, Helvetica, sans-serif;
      background: var(--bg);
      color: var(--text);
    }}
    main {{
      max-width: 1180px;
      margin: 0 auto;
      padding: 32px 20px 56px;
    }}
    header {{
      margin-bottom: 24px;
    }}
    h1 {{
      margin: 0 0 6px;
      font-size: 28px;
      letter-spacing: 0;
    }}
    p {{
      margin: 0;
      color: var(--muted);
    }}
    form.upload, section {{
      background: var(--panel);
      border: 1px solid var(--line);
      border-radius: 8px;
      padding: 20px;
      box-shadow: 0 1px 2px rgba(16, 24, 40, .04);
    }}
    form.upload {{
      display: grid;
      gap: 14px;
      margin-bottom: 20px;
    }}
    input[type=file] {{
      width: 100%;
      padding: 18px;
      border: 1px dashed #98a2b3;
      border-radius: 8px;
      background: #fbfcfe;
    }}
    button {{
      justify-self: start;
      border: 0;
      border-radius: 6px;
      padding: 10px 16px;
      color: #fff;
      background: var(--accent);
      font-weight: 700;
      cursor: pointer;
    }}
    button:hover {{ background: var(--accent-dark); }}
    .error {{
      margin: 0 0 16px;
      padding: 12px 14px;
      border: 1px solid #fecdca;
      border-radius: 8px;
      background: #fffbfa;
      color: var(--danger);
    }}
    .actions {{
      display: flex;
      justify-content: flex-end;
      margin-bottom: 14px;
    }}
    .tab-input {{
      position: absolute;
      opacity: 0;
      pointer-events: none;
    }}
    .tabs {{
      display: flex;
      gap: 8px;
      margin-bottom: 14px;
      border-bottom: 1px solid var(--line);
    }}
    .tabs label {{
      padding: 8px 12px;
      border: 1px solid transparent;
      border-bottom: 0;
      border-radius: 6px 6px 0 0;
      color: var(--muted);
      cursor: pointer;
      font-weight: 700;
    }}
    #tab-month:checked ~ .tabs label[for="tab-month"],
    #tab-day:checked ~ .tabs label[for="tab-day"] {{
      border-color: var(--line);
      background: #fff;
      color: var(--accent-dark);
    }}
    .tab-panel {{ display: none; }}
    #tab-month:checked ~ .month-panel,
    #tab-day:checked ~ .day-panel {{ display: block; }}
    table {{
      width: 100%;
      border-collapse: collapse;
      table-layout: fixed;
    }}
    th, td {{
      padding: 10px 12px;
      border-bottom: 1px solid var(--line);
      text-align: left;
      vertical-align: top;
      overflow-wrap: anywhere;
    }}
    th {{
      background: #eef3f7;
      color: #344054;
      font-size: 13px;
    }}
    .num {{ text-align: right; font-variant-numeric: tabular-nums; }}
    .warnings {{
      margin: 16px 0 0;
      color: #9a6700;
    }}
    @media (max-width: 760px) {{
      main {{ padding: 20px 12px 40px; }}
      table {{ display: block; overflow-x: auto; white-space: nowrap; }}
      th, td {{ min-width: 140px; }}
    }}
  </style>
</head>
<body>
  <main>
    <header>
      <h1>{APP_NAME}</h1>
      <p>TTK_Content: upload nhiều file Live/Video Excel để tính revenue_24h, content_impression và product_impression.</p>
    </header>
    {error_html}
    <form class="upload" method="post" enctype="multipart/form-data" action="/analyze">
      <input type="file" name="files" multiple accept=".xlsx">
      <button type="submit">Tính chỉ số</button>
    </form>
    {table_html}
  </main>
</body>
</html>""".encode("utf-8")


def parse_multipart(headers, body: bytes) -> list[tuple[str, bytes]]:
    content_type = headers.get("Content-Type", "")
    message = BytesParser(policy=default).parsebytes(
        f"Content-Type: {content_type}\r\nMIME-Version: 1.0\r\n\r\n".encode("utf-8") + body
    )
    files = []
    for part in message.iter_parts():
        filename = part.get_filename()
        if not filename:
            continue
        payload = part.get_payload(decode=True) or b""
        if payload:
            files.append((Path(filename).name, payload))
    return files


class AppHandler(BaseHTTPRequestHandler):
    server_version = "TTKContentTool/1.0"

    def do_GET(self):
        self.respond(200, render_page())

    def do_POST(self):
        try:
            if self.path == "/analyze":
                self.handle_analyze()
            elif self.path == "/export":
                self.handle_export()
            else:
                self.respond(404, b"Not found", "text/plain; charset=utf-8")
        except Exception as exc:  # pragma: no cover - visible local error
            traceback.print_exc()
            self.respond(500, render_page(f"Lỗi xử lý: {exc}"))

    def handle_analyze(self):
        length = int(self.headers.get("Content-Length", "0"))
        files = parse_multipart(self.headers, self.rfile.read(length))
        if not files:
            self.respond(400, render_page("Vui lòng chọn ít nhất một file .xlsx."))
            return

        parsed = []
        errors = []
        for filename, payload in files:
            if not filename.casefold().endswith(".xlsx"):
                errors.append(f"{filename}: chỉ hỗ trợ .xlsx")
                continue
            try:
                parsed.append(parse_workbook(payload, filename))
            except Exception as exc:
                errors.append(f"{filename}: {exc}")

        if not parsed:
            self.respond(400, render_page("; ".join(errors) or "Không đọc được file nào."))
            return

        duplicates = find_duplicate_files(parsed)
        if duplicates:
            self.respond(400, render_page("; ".join(errors + duplicates)))
            return

        rows = summarize(parsed)
        daily_rows = summarize_daily(parsed)
        error = "; ".join(errors)
        self.respond(200, render_page(error, rows, daily_rows))

    def handle_export(self):
        length = int(self.headers.get("Content-Length", "0"))
        body = self.rfile.read(length).decode("utf-8", errors="replace")
        payload_match = re.search(r"payload=([^&]+)", body)
        if not payload_match:
            self.respond(400, b"Missing payload", "text/plain; charset=utf-8")
            return
        from urllib.parse import unquote_plus

        data = json.loads(unquote_plus(payload_match.group(1)))
        if isinstance(data, list):
            monthly_data = data
            daily_data = []
        else:
            monthly_data = data.get("monthly", [])
            daily_data = data.get("daily", [])

        def rows_from_payload(items: list[dict[str, object]]) -> list[SummaryRow]:
            return [
                SummaryRow(
                    content_type=item["content_type"],
                    period=item["period"],
                    revenue_24h=float(item["revenue_24h"]),
                    content_impression=float(item["content_impression"]),
                    product_impression=float(item["product_impression"]),
                    files=list(item["files"]),
                    rows=int(item["rows"]),
                    warnings=list(item["warnings"]),
                )
                for item in items
            ]

        rows = rows_from_payload(monthly_data)
        daily_rows = rows_from_payload(daily_data)
        _, xlsx_path = save_outputs(rows, daily_rows)
        payload = xlsx_path.read_bytes()
        self.send_response(200)
        self.send_header(
            "Content-Type",
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        self.send_header("Content-Disposition", f'attachment; filename="{xlsx_path.name}"')
        self.send_header("Content-Length", str(len(payload)))
        self.end_headers()
        self.wfile.write(payload)

    def respond(self, status: int, body: bytes, content_type: str = "text/html; charset=utf-8"):
        self.send_response(status)
        self.send_header("Content-Type", content_type)
        self.send_header("Content-Length", str(len(body)))
        self.end_headers()
        self.wfile.write(body)

    def log_message(self, format, *args):
        print(f"[server] {self.address_string()} - {format % args}")


def cli(argv: list[str]) -> int:
    if len(argv) >= 2 and argv[1] == "--folder":
        folder = Path(argv[2]) if len(argv) >= 3 else Path("raw_sample")
        rows, daily_rows, parsed = parse_folder(folder)
        if not parsed:
            print(f"No .xlsx files found in {folder}")
            return 1
        duplicates = find_duplicate_files(parsed)
        if duplicates:
            print("Duplicate file error:")
            for duplicate in duplicates:
                print(f"- {duplicate}")
            return 1
        print("Live/Video,Thời gian,revenue_24h,product_impression,content_impression,Số file")
        for row in rows:
            print(
                f"{row.content_type},{row.period},{row.revenue_24h:.2f},"
                f"{row.product_impression:.0f},{row.content_impression:.0f},{len(row.files)}"
            )
        csv_path, xlsx_path = save_outputs(rows, daily_rows)
        print(f"Saved: {csv_path}")
        print(f"Saved: {xlsx_path}")
        return 0

    port = int(argv[1]) if len(argv) >= 2 else DEFAULT_PORT
    server = ThreadingHTTPServer(("127.0.0.1", port), AppHandler)
    print(f"{APP_NAME} is running at http://127.0.0.1:{port}")
    print("Press Ctrl+C to stop.")
    try:
        server.serve_forever()
    except KeyboardInterrupt:
        print("\nStopped.")
    finally:
        server.server_close()
    return 0


if __name__ == "__main__":
    raise SystemExit(cli(sys.argv))
